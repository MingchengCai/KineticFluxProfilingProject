"""
SP-KFP flux calculator
======================

GUI-based data-processing pipeline for **non-stationary isotope tracing /
kinetic flux profiling (SP-KFP)**.

Workflow
--------
    1. SUBMIT DATA
       An .xlsx workbook with two worksheets:
         * sheet 1: labeling fractions per isotopologue
         * sheet 2: pool sizes (one unit for all metabolites)
       Sheet 1: one header row containing 'isotopeLabel' and 'formula' (a
       'Metabolite' header and a 'Group' column are optional; without a
       'Metabolite' header the names are read from the column immediately left
       of isotopeLabel).  To the right of 'formula' come the sample columns
       '0_1', '0.05_2', ..., where the text before '_' is the time in hours and
       the text after '_' is the replicate number.
       Sheet 2: cell B1 = the unit, then column A = metabolite, column B = pool
       size.  A pool size is only needed for the metabolites used in the flux.

    2. FIT EVERY METABOLITE
       The labeled form of each metabolite is normalized by the mean labeling
       fraction of the last time point, and

           L(t) = 1 - exp(-k_app * t)

       is fitted with k_app as the only parameter (no Plateau).  k_app and the
       fitted curve are reported for every metabolite.

    3. CHOOSE THE FLUX INPUTS (dialog)
       The user picks exactly one pseudo tracer and one or more metabolites to
       pool for the flux, and confirms that these metabolites are downstream of
       the pseudo tracer (Criterion 3: major branch; Criterion 2: reversibility
       for the metabolites that are merged into one pool).

    4. CRITERION 1
       k_app(pseudo tracer) >= 4 x k_app of EVERY pooled metabolite; otherwise a
       dialog names the offending pairs and no flux is calculated.

    5. FLUX
       The labeled concentrations of the pooled metabolites are summed per
       sample and fitted with

           L(t) = Plateau * (1 - exp(-k_app * t))     flux = k_app * Plateau

       and its standard error is propagated from the two parameter errors,
       which are treated as independent.

Models
------
    per metabolite   L(t) = 1 - exp(-k_app t)                (normalized data)
    pooled/flux      L(t) = Plateau (1 - exp(-k_app t))      (concentrations)

Fitting
-------
Ordinary (unweighted) least squares on every individual sample, i.e.
replicates are treated as independent observations, as in R's nls().  The
per-metabolite model has one parameter, the pooled model two; both are solved
by scanning the residual sum of squares over log k (with the Plateau profiled
out analytically when it exists) and then refining by a golden-section search,
which is deterministic and free of starting-value dependence.  Standard errors
are the asymptotic (Gauss-Newton) errors (J'J)^-1 * RSS / dof.  k_app and
Plateau are then *treated as independent* when the flux uncertainty is
propagated, so no covariance term is used:

    SE(flux)^2 = Plateau^2 * SE(k_app)^2 + k_app^2 * SE(Plateau)^2

(The least-squares estimates of the two parameters are in fact strongly
correlated; the fitted covariance is still computed and kept in
FitResult.cov_kp / FluxResult.cov_kp, it is simply not used here.)

Usage
-----
    python spkfp_app.py                        # start the GUI
    python spkfp_app.py --check data.xlsx      # audit + fit (text report)
    python spkfp_app.py --check data.xlsx --pt IMP --flux AMP,ADP,ATP

Dependencies
------------
The program needs numpy, matplotlib and openpyxl.  They do not have to be
installed beforehand: if an import fails, the missing packages are installed
with the running interpreter (pip, PyPI or a mirror, `--user` if the
environment is not writable) and the program restarts itself.  Set the
environment variable `SPKFP_NO_AUTO_INSTALL=1` to switch that off, and
`SPKFP_PIP_INDEX_URL` to force a particular package index.
Only Python 3.8+ itself (with tkinter) has to be present.
"""


from __future__ import annotations

import importlib.util          # also binds "importlib" (invalidate_caches, import_module)
import math
import os
import re
import subprocess
import sys
import tempfile
import traceback
import urllib.request
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Sequence, Tuple

# ---------------------------------------------------------------------------
# Dependency bootstrap
#
# The program needs three third-party packages.  They are imported below; if one
# of them is missing (a fresh computer, a new virtual environment, a bare Python
# installation) the imports fail, and the bootstrap then installs the missing
# packages with the very interpreter that is running this file and restarts the
# program.  So a colleague can copy the .py (and the .xlsx) onto a machine that
# only has Python and simply run it.
#
# Nothing here needs a third-party module, so it is safe to run first.
# ---------------------------------------------------------------------------

#: (import name, requirement passed to pip)
REQUIRED_PACKAGES: Tuple[Tuple[str, str], ...] = (
    ("numpy", "numpy>=1.20"),
    ("matplotlib", "matplotlib>=3.4"),
    ("openpyxl", "openpyxl>=3.0"),
)

_DEP_FAILED: Dict[str, BaseException] = {}
_BOOTSTRAP_FLAG = "SPKFP_BOOTSTRAP_DONE"
_NO_INSTALL_ENV = "SPKFP_NO_AUTO_INSTALL"
def _env_int(name: str, default: int) -> int:
    try:
        value = int(float(os.environ.get(name, "")))
        return value if value > 0 else default
    except (TypeError, ValueError):
        return default


#: seconds a single pip download attempt may take before the next index is tried
_PIP_TIMEOUT = _env_int("SPKFP_PIP_TIMEOUT", 300)
_PROBE_TIMEOUT = 5          # seconds to wait for an index before skipping it
_INDEX_ENV = "SPKFP_PIP_INDEX_URL"     # a single index to try first
_INDEXES_ENV = "SPKFP_PIP_INDEXES"     # a complete, custom chain (; or , separated)
_SKIP_CN_ENV = "SPKFP_SKIP_CN_MIRRORS"  # never fall back to the Chinese mirrors

#: Chinese mirrors, only used as a fallback when PyPI itself does not answer
_CN_MIRRORS = (
    "https://pypi.tuna.tsinghua.edu.cn/simple",
    "https://mirrors.aliyun.com/pypi/simple/",
)


def _pip_indexes() -> Tuple[Optional[str], ...]:
    """The order in which package indexes are tried.

    PyPI (or whatever the machine's pip is configured for) is always first, so a
    user outside China normally never touches the Chinese mirrors; they are only
    a fallback for networks where pypi.org is slow or blocked.  A user can
    replace the whole chain with SPKFP_PIP_INDEXES, put one index first with
    SPKFP_PIP_INDEX_URL, or drop the Chinese mirrors with SPKFP_SKIP_CN_MIRRORS=1.
    """
    custom = os.environ.get(_INDEXES_ENV, "").strip()
    if custom:
        return tuple(x.strip() for x in re.split(r"[;,]", custom) if x.strip())
    chain: List[Optional[str]] = []
    first = os.environ.get(_INDEX_ENV, "").strip()
    if first:
        chain.append(first)          # an explicit choice wins over everything
    chain.append(None)               # pip default: PyPI, or the pip.conf of the machine
    if not os.environ.get(_SKIP_CN_ENV):
        chain.extend(_CN_MIRRORS)
    return tuple(chain)


def _log(message: str) -> None:
    """print() that survives pythonw.exe (where stdout may be missing)."""
    try:
        if sys.stdout is not None:
            print(message, flush=True)
    except Exception:
        pass


def _progress_window(message: str):
    """Show "please wait" while pip runs; returns a callable that closes it."""
    try:
        import tkinter as tk
        from tkinter import ttk

        root = tk.Tk()
        root.title("SP-KFP - installing required packages")
        root.resizable(False, False)
        try:
            root.attributes("-topmost", True)
        except Exception:
            pass
        frame = ttk.Frame(root, padding=18)
        frame.grid(row=0, column=0)
        ttk.Label(frame, text=message, justify="left", wraplength=520,
                  font=("Segoe UI", 10)).grid(row=0, column=0, sticky="w")
        ttk.Label(frame, text="This happens only once and may take a few minutes.",
                  foreground="#555555").grid(row=1, column=0, sticky="w", pady=(8, 0))
        bar = ttk.Progressbar(frame, mode="indeterminate", length=480)
        bar.grid(row=2, column=0, sticky="ew", pady=(12, 0))
        bar.start(40)
        root.update_idletasks()
        root.update()

        def close():
            try:
                bar.stop()
                root.destroy()
            except Exception:
                pass

        return close, root
    except Exception:
        return (lambda: None), None


def _fatal(title: str, message: str) -> None:
    """Tell the user that the automatic installation did not work."""
    _log("")
    _log("=" * 72)
    _log(title)
    _log(message)
    _log("=" * 72)
    try:
        import tkinter as tk
        from tkinter import messagebox

        root = tk.Tk()
        root.withdraw()
        try:
            root.attributes("-topmost", True)
        except Exception:
            pass
        messagebox.showerror(title, message)
        root.destroy()
    except Exception:
        pass


def _pip_available() -> bool:
    try:
        return importlib.util.find_spec("pip") is not None
    except Exception:
        return False


def _ensure_pip() -> bool:
    if _pip_available():
        return True
    _log("pip is missing - trying to enable it with ensurepip ...")
    try:
        subprocess.run([sys.executable, "-m", "ensurepip", "--upgrade"],
                       timeout=180, check=False,
                       **(_subprocess_flags()))
    except Exception:
        return False
    importlib.invalidate_caches()
    return _pip_available()


def _subprocess_flags() -> dict:
    """Hidden console window on Windows (the GUI already shows progress)."""
    if os.name == "nt":
        return {"creationflags": getattr(subprocess, "CREATE_NO_WINDOW", 0)}
    return {}


def _pip_install(requirements: Sequence[str], index_url: Optional[str],
                 user: bool) -> Tuple[int, str]:
    """Run pip once.  Returns (returncode, tail of its output)."""
    cmd = [sys.executable, "-m", "pip", "install", "--upgrade",
           "--disable-pip-version-check", "--no-input", "--prefer-binary",
           "--timeout", "20", "--retries", "1"]
    if user:
        cmd.append("--user")
    if index_url:
        cmd += ["--index-url", index_url]
    cmd += list(requirements)

    log_path = os.path.join(tempfile.gettempdir(), "spkfp_pip_install.log")
    live = bool(getattr(sys.stdout, "isatty", lambda: False)())
    try:
        if live:
            proc = subprocess.run(cmd, timeout=_PIP_TIMEOUT, check=False,
                                  **(_subprocess_flags()))
            tail = ""
        else:
            with open(log_path, "w", encoding="utf-8", errors="replace") as fh:
                proc = subprocess.run(cmd, stdout=fh, stderr=subprocess.STDOUT,
                                      timeout=_PIP_TIMEOUT, check=False,
                                      **(_subprocess_flags()))
            try:
                with open(log_path, encoding="utf-8", errors="replace") as fh:
                    tail = fh.read()[-1800:]
            except Exception:
                tail = ""
    except subprocess.TimeoutExpired:
        return 1, f"pip timed out after {_PIP_TIMEOUT} s"
    except Exception as exc:  # pip not runnable at all
        return 1, f"{type(exc).__name__}: {exc}"
    return proc.returncode, tail


def _index_alive(url: Optional[str]) -> bool:
    """Quick reachability probe for one package index (a dead index must not hang us)."""
    target = url or "https://pypi.org/simple/"
    if target.startswith("file:"):
        return True
    try:
        req = urllib.request.Request(target, method="HEAD",
                                     headers={"User-Agent": "pip/23"})
        with urllib.request.urlopen(req, timeout=_PROBE_TIMEOUT) as resp:
            return 200 <= resp.status < 400
    except Exception:
        return False


def _pip_configured_index() -> str:
    """The index-url that pip itself is configured with ('' when there is none)."""
    try:
        out = subprocess.run([sys.executable, "-m", "pip", "config", "get",
                              "global.index-url"],
                             capture_output=True, text=True, timeout=30,
                             check=False, **(_subprocess_flags()))
        return (out.stdout or "").strip()
    except Exception:
        return ""


def _install_requirements(requirements: Sequence[str]) -> Tuple[bool, str]:
    """Try several indexes / --user until one attempt succeeds."""
    if not _ensure_pip():
        return False, "pip is not available for this Python installation."

    seen = []
    hinted = False
    tail = ""            # used by the caller when every index was skipped
    for index in _pip_indexes():
        if index is not None and not index:
            continue
        if index in seen:
            continue
        seen.append(index)
        label = index or "PyPI (default)"
        if index is None:
            # "pip's default" may be PyPI or a configured mirror; only skip it when
            # nothing is configured and pypi.org does not answer at all
            if not _pip_configured_index() and not _index_alive(None):
                _log(f"skipping {label}: no answer from pypi.org within "
                     f"{_PROBE_TIMEOUT} s")
                continue
        elif not _index_alive(index):
            _log(f"skipping {label}: no answer within {_PROBE_TIMEOUT} s")
            continue
        _log(f"installing {', '.join(requirements)} from {label} ...")
        rc, tail = _pip_install(requirements, index, user=False)
        if rc == 0:
            return True, tail
        low = tail.lower()
        if "timed out" in low and not hinted:
            hinted = True
            _log("  an index that is too slow is skipped; set SPKFP_PIP_INDEX_URL "
                 "(or SPKFP_PIP_INDEXES) to use a faster mirror")
        if "permission" in low or "denied" in low or "not writable" in low:
            _log("  retrying for the current user only (--user) ...")
            rc, tail = _pip_install(requirements, index, user=True)
            if rc == 0:
                return True, tail
    return False, tail


def _pip_advice(requirements: Sequence[str]) -> str:
    # quote the specifiers: "pip install numpy>=1.20" would be parsed as a shell
    # redirection on Windows
    quoted = [f'"{r}"' if ("<" in r or ">" in r) else r for r in requirements]
    return f'    "{sys.executable}" -m pip install ' + " ".join(quoted)


def _reexec() -> bool:
    """Restart this script with the same interpreter and arguments."""
    if os.environ.get(_BOOTSTRAP_FLAG):
        return False
    try:
        script = os.path.abspath(__file__)
    except NameError:
        return False
    if not os.path.isfile(script):
        return False
    env = dict(os.environ)
    env[_BOOTSTRAP_FLAG] = "1"
    args = [sys.executable, script] + sys.argv[1:]
    _log("restarting ...")
    # On Windows os.execv passes the argument list to the C runtime without
    # quoting, so a path containing spaces would be split into several
    # arguments.  Start a fresh process instead (Popen quotes correctly) and let
    # this one finish.
    try:
        if os.name == "nt":
            subprocess.Popen(args, env=env, close_fds=True)
            raise SystemExit(0)
        os.execve(sys.executable, args, env)
    except SystemExit:
        raise
    except Exception:
        try:
            subprocess.Popen(args, env=env)
        except Exception:
            return False
        raise SystemExit(0)
    return True  # not reached


def _rebind_modules(missing: Sequence[Tuple[str, str]]) -> bool:
    """Last resort when the process cannot be restarted: import into globals."""
    ok = True
    for mod, _req in missing:
        try:
            if mod == "matplotlib":
                figure_mod = importlib.import_module("matplotlib.figure")
                globals()["Figure"] = getattr(figure_mod, "Figure")
            elif mod == "numpy":
                globals()["np"] = importlib.import_module("numpy")
            elif mod == "openpyxl":
                globals()["openpyxl"] = importlib.import_module("openpyxl")
        except Exception:
            ok = False
    return ok


def _bootstrap_failed(missing: Sequence[Tuple[str, str]], detail: str) -> None:
    names = ", ".join(mod for mod, _ in missing)
    reason = (detail.strip().splitlines()[-1] if detail.strip()
              else "no package index could be reached within 5 s - is there an "
                   "internet connection?")
    _fatal(
        "SP-KFP - required packages could not be installed",
        "This program needs the Python packages:\n"
        f"    {names}\n\n"
        "They are not installed here and the automatic installation failed.\n\n"
        f"Reason: {reason}\n\n"
        "Please install them once by hand, in a command prompt, with:\n"
        f"{_pip_advice([req for _mod, req in missing])}\n\n"
        "If this computer has no internet access, use an offline wheel package or "
        "run the program on a machine where the packages are available.  After the "
        "installation, start this program again.",
    )
    raise SystemExit(1)


def _install_and_restart(missing: Sequence[Tuple[str, str]]) -> None:
    """Install the missing packages, then restart (or rebind) so they are usable."""
    requirements = [req for _mod, req in missing]
    names = ", ".join(mod for mod, _ in missing)

    if os.environ.get(_NO_INSTALL_ENV):
        _bootstrap_failed(
            missing,
            f"automatic installation is disabled by {_NO_INSTALL_ENV}",
        )

    _log(f"missing Python package(s): {names} - installing them now")
    close, root = _progress_window(
        "The following Python packages are missing and will be installed now:\n"
        f"    {names}\n\n"
        "Please wait, the program continues by itself when this is done."
    )
    try:
        ok, tail = _install_requirements(requirements)
    finally:
        # keep the window responsive for one more moment, then close it
        try:
            if root is not None:
                root.update_idletasks()
                root.update()
        except Exception:
            pass
        close()

    if not ok:
        _bootstrap_failed(missing, tail)

    importlib.invalidate_caches()
    _log("installation finished")

    if _reexec():
        raise SystemExit(0)

    if not _rebind_modules(missing):
        _bootstrap_failed(
            missing,
            "the packages were installed but this process could not reload them; "
            "please start the program again",
        )


# ---------------------------------------------------------------------------
# Third-party imports (guarded so that a missing package can be installed)
# ---------------------------------------------------------------------------

try:  # pragma: no cover - environment dependent
    import numpy as np
except Exception as _exc:  # pragma: no cover
    np = None  # type: ignore[assignment]
    _DEP_FAILED["numpy"] = _exc

try:  # pragma: no cover - environment dependent
    import openpyxl
except Exception as _exc:  # pragma: no cover
    openpyxl = None  # type: ignore[assignment]
    _DEP_FAILED["openpyxl"] = _exc

try:  # pragma: no cover - environment dependent
    from matplotlib.figure import Figure

    _MPL_ERR = None
except Exception as _exc:  # pragma: no cover
    Figure = None  # type: ignore[assignment]
    _MPL_ERR = str(_exc)
    _DEP_FAILED["matplotlib"] = _exc

try:  # pragma: no cover - environment dependent
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox
except Exception:  # pragma: no cover
    tk = None
    ttk = None

if _DEP_FAILED:
    _install_and_restart(
        [(mod, req) for mod, req in REQUIRED_PACKAGES if mod in _DEP_FAILED]
    )


# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

PARENT_TOKEN = "PARENT"
HEADER_KEYS = ("group", "metabolite", "isotopelabel", "formula")
TIME_RE = re.compile(r"^\s*(\d+(?:\.\d+)?)\s*_\s*(\d+)\s*$")
FORMULA_RE = re.compile(r"([A-Z][a-z]?)(\d*)")
LABEL_RE = re.compile(
    r"([A-Za-z]{1,2})\s*(\d{1,3})\s*[-_ ]?\s*label\s*[-_ ]?(\d{1,2})", re.I
)
LABEL_ONLY_RE = re.compile(r"(?:label|m)\s*[-+ ]?\s*(\d{1,2})", re.I)

CRITERION1_FACTOR = 4.0      # k_app(pseudo tracer) >= 4 x k_app(flux metabolites)
FIT_K_MIN_FACTOR = 1e-5      # k search range: [1e-5/tmax, 1e3/tmin]
FIT_K_MAX_FACTOR = 1e3

MINOR = 1e-6


# ---------------------------------------------------------------------------
# Errors
# ---------------------------------------------------------------------------


class ValidationError(Exception):
    """Raised when the submitted workbook does not meet the requirements."""

    def __init__(self, problems: Sequence[str], title: str = "Data file rejected"):
        self.problems = [str(p) for p in problems]
        self.title = title
        super().__init__("\n".join(self.problems))


class PipelineError(Exception):
    """Raised when the data pass the audit but the analysis cannot proceed."""


# ---------------------------------------------------------------------------
# Small utilities
# ---------------------------------------------------------------------------


def _s(value) -> str:
    """Cell -> stripped string ('' for empty)."""
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    return str(value).strip()


def _norm_header(value) -> str:
    """Header text -> comparison key ('isotopeLabel' -> 'isotopelabel')."""
    return re.sub(r"[^a-z0-9]", "", _s(value).lower())


def _to_float(value) -> Optional[float]:
    """Cell -> float, or None when it is not a number."""
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if isinstance(value, float) and math.isnan(value):
            return None
        return float(value)
    txt = _s(value)
    if not txt:
        return None
    txt = txt.replace(",", "")
    try:
        return float(txt)
    except ValueError:
        return None


def num(value, sig: int = 4, na: str = "n/a") -> str:
    """Format a number with `sig` significant digits (ASCII safe)."""
    try:
        x = float(value)
    except (TypeError, ValueError):
        return na
    if not math.isfinite(x):
        return na
    if x == 0:
        return "0"
    return f"{x:.{sig}g}"


def parse_formula(formula: str) -> Dict[str, int]:
    """'C10H13N4O8P' -> {'C': 10, 'H': 13, 'N': 4, 'O': 8, 'P': 1}."""
    out: Dict[str, int] = {}
    for elem, cnt in FORMULA_RE.findall(_s(formula)):
        out[elem] = out.get(elem, 0) + (int(cnt) if cnt else 1)
    return out


def parse_label(label: str) -> Tuple[Optional[str], Optional[int]]:
    """'N15-label-4' -> ('N', 4).  ('M+4' -> (None, 4))."""
    txt = _s(label)
    m = LABEL_RE.search(txt)
    if m:
        return m.group(1).capitalize() if len(m.group(1)) == 1 else m.group(1), int(m.group(3))
    m = LABEL_ONLY_RE.search(txt)
    if m:
        return None, int(m.group(1))
    return None, None


def mean_sd_by_time(t: np.ndarray, y: np.ndarray) -> Tuple[np.ndarray, np.ndarray, np.ndarray]:
    """Mean and sample SD of y at every distinct time point."""
    ut = np.unique(np.asarray(t, dtype=float))
    means = np.empty(ut.size)
    sds = np.empty(ut.size)
    for i, tv in enumerate(ut):
        sel = np.asarray(t, dtype=float) == tv
        vals = np.asarray(y, dtype=float)[sel]
        vals = vals[np.isfinite(vals)]
        means[i] = vals.mean() if vals.size else np.nan
        sds[i] = vals.std(ddof=1) if vals.size > 1 else 0.0
    return ut, means, sds


# ---------------------------------------------------------------------------
# Data model
# ---------------------------------------------------------------------------


@dataclass
class SampleColumn:
    key: str
    time: float
    sample: int
    col: int = 0          # 0-based worksheet column index of the raw data

    @property
    def label(self) -> str:
        return f"{self.time:g} h #{self.sample}"


@dataclass
class Metabolite:
    name: str
    group: str
    formula: str
    labels: Dict[str, np.ndarray]       # labeled (non-PARENT) forms -> fractions
    parent_labels: List[str]
    peak_label: str                     # text shown in the results table
    weight_note: str
    poolsize: Optional[float] = None
    order: int = 0

    @property
    def labeled_forms(self) -> List[str]:
        return list(self.labels)

    def series(self, n_cols: int) -> np.ndarray:
        """Labeling value per sample column (fraction, or conc. if poolsize)."""
        arr = np.zeros(n_cols, dtype=float)
        for lab in self.labels:
            arr = arr + _weighted_form(self, lab, n_cols)
        return arr

    def concentration(self, n_cols: int) -> np.ndarray:
        """Labeled concentration per sample column (same units as sheet 2)."""
        base = self.series(n_cols)
        if self.poolsize is None:
            return base
        return base * float(self.poolsize)


def _weighted_form(met: "Metabolite", lab: str, n_cols: int) -> np.ndarray:
    """Contribution of one labeled form to the representative labeling value.

    With a single labeled form the raw fraction is used (no scaling).  With
    several forms the SOP weighting is applied:
        value = SUM(fraction x labeled atoms) / (atoms of that element)
    """
    vals = np.asarray(met.labels[lab], dtype=float)
    if len(met.labels) == 1:
        return vals
    elem, n_atoms = parse_label(lab)
    n_elem = parse_formula(met.formula).get(elem) if elem else None
    if n_atoms and n_elem and n_elem >= n_atoms:
        return vals * (float(n_atoms) / float(n_elem))
    return vals / float(len(met.labels))


@dataclass
class Dataset:
    path: str
    columns: List[SampleColumn]
    metabolites: List[Metabolite]
    pool_unit: str
    notes: List[str] = field(default_factory=list)

    @property
    def times(self) -> np.ndarray:
        return np.array([c.time for c in self.columns], dtype=float)

    @property
    def n_cols(self) -> int:
        return len(self.columns)

    def by_group(self, group: str) -> List[Metabolite]:
        """Group column is informational only; kept for convenience."""
        return [m for m in self.metabolites if m.group == group]

    def unique_times(self) -> List[float]:
        return sorted({c.time for c in self.columns})

    def metabolite(self, name: str) -> Optional[Metabolite]:
        for m in self.metabolites:
            if m.name == name:
                return m
        return None

    def with_pool(self) -> List[str]:
        return [m.name for m in self.metabolites if m.poolsize is not None]

    def summary(self) -> str:
        ts = self.unique_times()
        tspan = f"{ts[0]:g}-{ts[-1]:g} h" if ts else "?"
        with_pool = len(self.with_pool())
        return (
            f"{len(self.metabolites)} metabolites ({with_pool} with a pool size); "
            f"{len(ts)} time points ({tspan}); {self.n_cols} samples; "
            f"pool unit: {self.pool_unit}"
        )


@dataclass
class FitResult:
    name: str
    group: str
    label: str
    n: int = 0
    k: float = float("nan")
    k_se: float = float("nan")
    plateau: float = float("nan")
    plateau_se: float = float("nan")
    cov_kp: float = float("nan")
    r2: float = float("nan")
    unit: str = ""
    ok: bool = False
    message: str = ""
    warn: str = ""
    model: str = "normalized"          # "normalized" (per metabolite) | "saturation" (pooled)
    rate: float = float("nan")         # k*Plateau of the pooled fit
    rate_se: float = float("nan")
    tmax: float = float("nan")
    norm: float = float("nan")         # last-time-point mean used for normalization
    instantaneous: bool = False        # k_app hit the upper search bound
    # plotting helpers
    t: Optional[np.ndarray] = None
    y: Optional[np.ndarray] = None

    @property
    def key(self) -> str:
        return f"{self.name}|{self.label}"


@dataclass
class Criterion1Result:
    evaluated: bool = False
    passed: bool = True
    lines: List[str] = field(default_factory=list)
    failures: List[str] = field(default_factory=list)


@dataclass
class FluxResult:
    """Pooled flux: flux = k_app * Plateau on the summed labeled concentrations."""
    pt: str
    members: List[str]
    unit: str
    flux: float = float("nan")
    flux_se: float = float("nan")
    k: float = float("nan")
    k_se: float = float("nan")
    plateau: float = float("nan")
    plateau_se: float = float("nan")
    cov_kp: float = float("nan")
    r2: float = float("nan")
    n: int = 0
    note: str = ""


# ---------------------------------------------------------------------------
# 1. Load + audit the workbook
# ---------------------------------------------------------------------------


def _find_header(rows: List[List[object]]) -> Tuple[Optional[int], Dict[str, int], str]:
    """Locate the header row carrying 'isotopeLabel' and 'formula'.

    'Metabolite' and 'Group' headers are optional: the metabolite name may sit in
    a column without a header (immediately left of 'isotopeLabel', as in the
    template) and the group column no longer exists.
    """
    best_idx, best_found, best_n = None, {}, 0
    for i, row in enumerate(rows[:60]):
        found: Dict[str, int] = {}
        for j, cell in enumerate(row):
            key = _norm_header(cell)
            if key in HEADER_KEYS and key not in found:
                found[key] = j
        if len(found) > best_n:
            best_idx, best_found, best_n = i, found, len(found)
        if "isotopelabel" in found and "formula" in found:
            return i, found, ""
    readable = {
        "group": "Group",
        "metabolite": "Metabolite",
        "isotopelabel": "isotopeLabel",
        "formula": "formula",
    }
    if best_idx is None:
        return None, {}, (
            "Sheet 1 is empty: no header row containing 'isotopeLabel' and 'formula' "
            "was found."
        )
    missing = [k for k in ("isotopelabel", "formula") if k not in best_found]
    return None, best_found, (
        "Sheet 1: the header row must contain 'isotopeLabel' and 'formula'. "
        f"The most complete row found (row {best_idx + 1}) contains "
        f"{', '.join(readable[k] for k in best_found) or 'nothing recognizable'} but "
        f"is missing {', '.join(readable[k] for k in missing)}."
    )


def load_dataset(path: str) -> Dataset:
    """Read and audit an SP-KFP workbook.  Raises ValidationError with reasons."""
    if openpyxl is None:
        raise RuntimeError("openpyxl is required to read .xlsx files (pip install openpyxl)")

    problems: List[str] = []
    notes: List[str] = []

    try:
        try:
            wb = openpyxl.load_workbook(path, data_only=True, rich_text=True)
        except TypeError:  # openpyxl < 3.1 has no rich_text switch
            wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        raise ValidationError([f"The workbook could not be opened: {exc}"])
    if wb is None:
        raise ValidationError(["The workbook could not be opened."])

    if len(wb.worksheets) < 2:
        raise ValidationError(
            [
                f"The workbook contains {len(wb.worksheets)} worksheet(s). "
                "Two worksheets are required: sheet 1 with the labeling fractions "
                "and sheet 2 with the pool sizes."
            ]
        )

    ws1, ws2 = wb.worksheets[0], wb.worksheets[1]
    rows1 = [list(r) for r in ws1.iter_rows(values_only=True)]
    rows2 = [list(r) for r in ws2.iter_rows(values_only=True)]

    # ---------------- header ----------------
    header_idx, colmap, header_msg = _find_header(rows1)
    if header_idx is None:
        raise ValidationError([header_msg])

    header_row = rows1[header_idx]

    # column map: isotopeLabel and formula are mandatory; the metabolite column is
    # either explicitly labelled or sits immediately left of isotopeLabel
    col_lab = colmap["isotopelabel"]
    col_form = colmap["formula"]
    col_met = colmap.get("metabolite")
    met_header_found = col_met is not None
    if col_met is None:
        col_met = col_lab - 1 if col_lab > 0 else 0
    col_group = colmap.get("group")      # optional, display only

    n_cols_header = len(header_row)
    data_start = col_form + 1

    # ---------------- sample columns ----------------
    columns: List[SampleColumn] = []
    bad_headers: List[str] = []
    seen_keys: Dict[Tuple[float, int], str] = {}
    for j in range(data_start, n_cols_header):
        raw = header_row[j]
        txt = _s(raw)
        if not txt:
            continue
        m = TIME_RE.match(txt)
        if not m:
            bad_headers.append(txt)
            continue
        t_val, s_val = float(m.group(1)), int(m.group(2))
        if (t_val, s_val) in seen_keys:
            problems.append(
                f"Sheet 1: the sample column '{txt}' is duplicated "
                f"(time {t_val:g} h, sample {s_val})."
            )
            continue
        seen_keys[(t_val, s_val)] = txt
        columns.append(SampleColumn(key=txt, time=t_val, sample=s_val, col=j))

    if bad_headers:
        problems.append(
            "Sheet 1: every column to the right of 'formula' must be a sample column "
            "named '<time>_<sample>' (e.g. '0.5_2', time in hours). These headers do "
            "not follow that pattern: " + ", ".join(repr(h) for h in bad_headers[:8])
        )
    if len(columns) < 2:
        problems.append(
            f"Sheet 1: only {len(columns)} valid sample column(s) were found; at least "
            "two sample columns are needed."
        )
    unique_times = sorted({c.time for c in columns})
    if len(columns) >= 2 and len(unique_times) < 2:
        problems.append(
            "Sheet 1: all sample columns belong to a single time point; at least two "
            "different time points are required."
        )

    # ---------------- data rows ----------------
    n_cols = len(columns)
    data_cols = [c.col for c in columns]

    records: List[Tuple[str, str, str, str, int, Dict[str, float]]] = []
    blank_cells = 0

    for i in range(header_idx + 1, len(rows1)):
        row = rows1[i]
        excel_row = i + 1
        if not any(_s(c) for c in row):
            continue
        g_raw = row[col_group] if (col_group is not None and col_group < len(row)) else None
        m_raw = row[col_met] if col_met < len(row) else None
        l_raw = row[col_lab] if col_lab < len(row) else None
        f_raw = row[col_form] if col_form < len(row) else None

        name, lab = _s(m_raw), _s(l_raw)
        has_values = any(
            j < len(row) and _s(row[j]) not in ("",) for j in data_cols
        )

        if not name and not lab:
            if has_values:
                problems.append(
                    f"Sheet 1, row {excel_row}: labeling values are present but the "
                    "metabolite name and 'isotopeLabel' are empty. Every data row must "
                    "carry the metabolite name and the isotope label together."
                )
            continue
        if not name:
            problems.append(
                f"Sheet 1, row {excel_row}: the metabolite name is empty although an "
                f"isotope label ({lab}) is given."
            )
            continue
        if not lab:
            problems.append(
                f"Sheet 1, row {excel_row} (metabolite '{name}'): 'isotopeLabel' is "
                "empty."
            )
            continue

        group = _s(g_raw)     # optional, informational only

        values: Dict[str, float] = {}
        for key, j in [(c.key, c.col) for c in columns]:
            if j >= len(row):
                values[key] = float("nan")
                blank_cells += 1
                continue
            val = _to_float(row[j])
            if val is None:
                if _s(row[j]) == "":
                    blank_cells += 1
                    values[key] = float("nan")
                else:
                    problems.append(
                        f"Sheet 1, row {excel_row} (metabolite '{name}', label '{lab}'): "
                        f"the value in column '{key}' is not a number ({_s(row[j])!r})."
                    )
                continue
            if not (val >= -MINOR and val <= 1.0 + 1e-6):
                problems.append(
                    f"Sheet 1, row {excel_row} (metabolite '{name}', label '{lab}'): "
                    f"labeling fraction {val:g} in column '{key}' is outside 0-1. "
                    "Labeling fractions must be proportions; if your data are "
                    "percentages, divide them by 100."
                )
            values[key] = val
        records.append((group, name, lab, _s(f_raw), excel_row, values))

    if blank_cells:
        notes.append(
            f"{blank_cells} empty sample cell(s) were found; those points are excluded "
            "from the fits."
        )

    # ---------------- assemble metabolites ----------------
    order: List[str] = []
    by_name: Dict[str, Dict[str, object]] = {}
    for group, name, lab, formula, excel_row, values in records:
        if name not in by_name:
            by_name[name] = {
                "group": group,
                "formula": formula,
                "rows": {},
                "excel_row": excel_row,
            }
            order.append(name)
        entry = by_name[name]
        if entry["group"] and group and entry["group"] != group:
            notes.append(
                f"'{name}' appears with more than one group label "
                f"({entry['group']!r} and {group!r}); the group column is only "
                "informational and was ignored."
            )
        if formula and not entry["formula"]:
            entry["formula"] = formula
        rows = entry["rows"]  # type: ignore[assignment]
        if lab in rows:
            problems.append(
                f"Sheet 1: metabolite '{name}' has the isotope label '{lab}' twice "
                f"(second occurrence in row {excel_row})."
            )
            continue
        rows[lab] = (values, excel_row)  # type: ignore[index]

    metabolites: List[Metabolite] = []
    for idx, name in enumerate(order):
        entry = by_name[name]
        group = str(entry["group"])
        formula = str(entry["formula"])
        rows = entry["rows"]  # type: ignore[assignment]
        parent_labels = [l for l in rows if PARENT_TOKEN.lower() in l.lower()]
        labeled = {l: rows[l] for l in rows if l not in parent_labels}
        if not labeled:
            problems.append(
                f"Sheet 1: metabolite '{name}' has no labeled form; at least one row "
                "without 'PARENT' in 'isotopeLabel' is required."
            )
            continue

        labels: Dict[str, np.ndarray] = {}
        for lab, (values, _row) in labeled.items():
            labels[lab] = np.array([values.get(c.key, np.nan) for c in columns], dtype=float)

        weight_note = ""
        if len(labels) > 1:
            parsed = [(l,) + parse_label(l) for l in labels]
            elems = {p[1] for p in parsed}
            atoms = [p[2] for p in parsed]
            counts = parse_formula(formula)
            usable = (
                len(elems) == 1
                and None not in elems
                and all(a is not None for a in atoms)
                and elems.pop() in counts
                and counts[parsed[0][1]] >= max(a for a in atoms if a)
            )
            if usable:
                elem = parsed[0][1]
                n_tot = counts[elem]
                weight_note = (
                    f"weighted average of {len(labels)} forms: "
                    f"SUM(fraction x labeled atoms)/{n_tot} {elem} atoms"
                )
                notes.append(
                    f"'{name}' has {len(labels)} labeled forms "
                    f"({', '.join(labels)}); they were averaged with the "
                    f"element-abundance weighting SUM(fraction x labeled atoms)/"
                    f"{n_tot} {elem} atoms."
                )
            else:
                weight_note = f"simple mean of {len(labels)} labeled forms"
                notes.append(
                    f"'{name}' has {len(labels)} labeled forms and the labeled element / "
                    "atom count could not be resolved from the labels or the formula; "
                    "a simple mean of the forms was used. Please supply a single "
                    "labeled form per metabolite."
                )

        if len(labels) == 1:
            only = list(labels)[0]
            peak_label = only
        else:
            peak_label = " + ".join(labels) + " (averaged)"

        metabolites.append(
            Metabolite(
                name=name,
                group=group,
                formula=formula,
                labels=labels,
                parent_labels=parent_labels,
                peak_label=peak_label,
                weight_note=weight_note,
                order=idx,
            )
        )

    # ---------------- sheet 2: pool sizes ----------------
    pool_unit = ""
    pool: Dict[str, float] = {}
    unit_row = None
    for i, row in enumerate(rows2[:3]):
        if len(row) > 1 and _s(row[1]) and _to_float(row[1]) is None:
            pool_unit = _s(row[1])
            unit_row = i
            break
    if unit_row is None:
        unit_row = -1
        notes.append(
            "Sheet 2: no unit was found in cell B1 of the pool-size sheet; the flux is "
            "reported in 'arbitrary units/h'."
        )
        pool_unit = "arbitrary unit"

    bad_pool: List[str] = []
    for i in range(unit_row + 1, len(rows2)):
        row = rows2[i]
        name = _s(row[0]) if len(row) > 0 else ""
        if not name:
            continue
        val = _to_float(row[1]) if len(row) > 1 else None
        if val is None:
            bad_pool.append(f"{name} (row {i + 1})")
            continue
        if val < 0:
            problems.append(
                f"Sheet 2: the pool size of '{name}' is negative ({val:g})."
            )
            continue
        if name in pool:
            problems.append(
                f"Sheet 2: the metabolite '{name}' appears more than once in the "
                "pool-size sheet."
            )
            continue
        pool[name] = val
    if bad_pool:
        problems.append(
            "Sheet 2: a numeric pool size is missing for " + ", ".join(bad_pool[:10]) + "."
        )

    missing = [m.name for m in metabolites if m.name not in pool]
    if missing:
        notes.append(
            "No pool size is given for these metabolite(s): " + ", ".join(missing)
            + ". They can still be fitted, but they cannot be used in the flux "
            "calculation (a labeled concentration needs a pool size)."
        )

    extra = [n for n in pool if n not in {m.name for m in metabolites}]
    if extra:
        notes.append(
            "Sheet 2 lists metabolite(s) that do not occur in sheet 1 and were "
            "ignored: " + ", ".join(extra) + "."
        )

    for met in metabolites:
        if met.name in pool:
            met.poolsize = pool[met.name]

    if problems:
        raise ValidationError(problems)

    dataset = Dataset(
        path=path,
        columns=columns,
        metabolites=metabolites,
        pool_unit=pool_unit,
        notes=notes,
    )
    return dataset


# ---------------------------------------------------------------------------
# 2. Single-exponential fitting
# ---------------------------------------------------------------------------


def _profile_ssr(k: float, t: np.ndarray, y: np.ndarray) -> Tuple[float, float]:
    """RSS and the optimal (analytically profiled) Plateau for a given k."""
    g = -np.expm1(-k * t)
    denom = float(g @ g)
    if denom <= 0:
        return float("inf"), 0.0
    p = float(g @ y) / denom
    r = y - p * g
    return float(r @ r), p


def _golden_min(f, a: float, b: float, iters: int = 200) -> float:
    gr = (math.sqrt(5.0) - 1.0) / 2.0
    c = b - gr * (b - a)
    d = a + gr * (b - a)
    fc, fd = f(c), f(d)
    for _ in range(iters):
        if fc < fd:
            b, d, fd = d, c, fc
            c = b - gr * (b - a)
            fc = f(c)
        else:
            a, c, fc = c, d, fd
            d = a + gr * (b - a)
            fd = f(d)
    return 0.5 * (a + b)


def fit_kapp(t, y) -> FitResult:
    """Fit L(t) = 1 - exp(-k_app * t), one parameter, no Plateau.

    `y` must already be normalized by the mean labeling fraction of the LAST
    time point, so that the fitted curve tends to 1 at the end of the window.
    """
    t = np.asarray(t, dtype=float).ravel()
    y = np.asarray(y, dtype=float).ravel()
    good = np.isfinite(t) & np.isfinite(y)
    t, y = t[good], y[good]
    res = FitResult(name="", group="", label="", n=int(t.size), t=t, y=y)

    n = t.size
    if n < 3:
        res.message = f"only {n} usable data point(s); at least 3 are required"
        return res
    t_pos = t[t > 0]
    if t_pos.size == 0 or float(t.max()) <= 0:
        res.message = "all time points are 0 h; the kinetics cannot be fitted"
        return res
    if not np.all(np.isfinite(y)):
        res.message = "the normalized labeling contains non-finite values"
        return res
    if float(np.ptp(y)) <= 1e-12:
        res.message = (
            "the normalized labeling is constant, so no rising curve can be fitted"
        )
        return res

    tmax = float(t.max())
    tmin = float(t_pos.min())
    res.tmax = tmax
    k_lo = FIT_K_MIN_FACTOR / tmax
    k_hi = FIT_K_MAX_FACTOR / tmin

    def ssr_of(logk: float) -> float:
        k = math.exp(logk)
        r = y - (-np.expm1(-k * t))
        return float(r @ r)

    ks = np.logspace(math.log10(k_lo), math.log10(k_hi), 500)
    g = -np.expm1(-np.outer(ks, t))
    r = y[None, :] - g
    ssr = np.einsum("ij,ij->i", r, r)
    i_best = int(np.argmin(ssr))

    if i_best == 0:
        # the curve wants to be flatter than the search floor allows
        res.message = (
            "k_app is too small to be identified: the normalized labeling never rises "
            "inside the sampled window"
        )
        return res

    lo = math.log(ks[max(i_best - 2, 0)])
    hi = math.log(ks[min(i_best + 2, ks.size - 1)])
    logk = _golden_min(ssr_of, lo, hi) if hi > lo else math.log(ks[i_best])
    k = math.exp(logk)

    model = -np.expm1(-k * t)
    r1 = y - model
    rss = float(r1 @ r1)
    tss = float(((y - y.mean()) ** 2).sum())
    res.r2 = 1.0 - rss / tss if tss > 0 else float("nan")

    dof = n - 1
    if dof >= 1:
        jac = t * np.exp(-k * t)          # dL/dk
        sxx = float(jac @ jac)
        if sxx > 0:
            res.k_se = float(math.sqrt(rss / dof / sxx))

    res.k = float(k)
    res.model = "normalized"
    res.ok = True
    # labeling complete before the first sample: k*t_min large means the curve is
    # flat at 1 everywhere and k_app is only a lower bound
    if k >= k_hi * 0.999 or k * tmin >= 5.0:
        res.instantaneous = True
        res.message = (
            "labeling is already complete at the first sampled time point, so k_app is "
            "not identifiable (it is larger than ~5/t_min)"
        )
        res.warn = (
            "labeling complete at the first time point: k_app is only a lower bound"
        )
        res.k_se = float("nan")
    if math.isfinite(res.r2) and res.r2 < 0.5 and not res.warn:
        res.warn = f"poor fit (R^2 = {res.r2:.2f})"
    return res


def fit_exponential(t, y) -> FitResult:
    """Fit L = Plateau * (1 - exp(-k_app t)) by least squares on all samples.

    This two-parameter form is used for the pooled labeled concentration in the
    flux step (flux = k_app * Plateau).  It is not used for the per-metabolite
    k_app, which is fitted on data normalized to the last time point.
    """
    t = np.asarray(t, dtype=float).ravel()
    y = np.asarray(y, dtype=float).ravel()
    good = np.isfinite(t) & np.isfinite(y)
    t, y = t[good], y[good]
    res = FitResult(name="", group="", label="", n=int(t.size), t=t, y=y)

    n = t.size
    if n < 3:
        res.message = f"only {n} usable data point(s); at least 3 are required"
        return res
    t_pos = t[t > 0]
    if t_pos.size == 0 or float(t.max()) <= 0:
        res.message = "all time points are 0 h; the kinetics cannot be fitted"
        return res
    if not np.any(np.abs(y) > 0):
        res.message = "no labeling was detected (every value is 0)"
        return res
    if float(np.ptp(y)) <= 0:
        res.message = (
            "the labeling values are constant, so a rising saturation curve cannot be "
            "fitted"
        )
        return res

    tmax = float(t.max())
    tmin = float(t_pos.min())
    res.tmax = tmax
    k_lo = FIT_K_MIN_FACTOR / tmax
    k_hi = FIT_K_MAX_FACTOR / tmin

    ks = np.logspace(math.log10(k_lo), math.log10(k_hi), 600)
    g = -np.expm1(-np.outer(ks, t))
    denom = np.einsum("ij,ij->i", g, g)
    p_of_k = (g @ y) / denom
    r = y[None, :] - p_of_k[:, None] * g
    ssr = np.einsum("ij,ij->i", r, r)
    i_best = int(np.argmin(ssr))

    if i_best == 0:
        res.message = (
            "k_app and Plateau are not separable for the pooled data: the labeling is "
            "still linear at the last time point, so the exponential fit cannot "
            "determine a Plateau (the pooled window is too short)"
        )
        return res

    lo = math.log(ks[max(i_best - 2, 0)])
    hi = math.log(ks[min(i_best + 2, ks.size - 1)])
    if hi > lo:
        logk = _golden_min(lambda lk: _profile_ssr(math.exp(lk), t, y)[0], lo, hi)
    else:
        logk = math.log(ks[i_best])
    k = math.exp(logk)
    _, plateau = _profile_ssr(k, t, y)

    if plateau <= 0:
        res.message = (
            "the labeling values do not show a saturating rise (fitted Plateau <= 0)"
        )
        return res
    if k >= k_hi * 0.999:
        res.message = (
            "the pooled labeling is already at its plateau at the earliest time point, "
            "so k_app is larger than the sampling can resolve"
        )
        return res

    g1 = -np.expm1(-k * t)
    r1 = y - plateau * g1
    rss = float(r1 @ r1)
    tss = float(((y - y.mean()) ** 2).sum())
    res.r2 = 1.0 - rss / tss if tss > 0 else float("nan")

    dof = n - 2
    if dof >= 1:
        e = np.exp(-k * t)
        J = np.column_stack([g1, plateau * t * e])   # d/dPlateau, d/dk
        JtJ = J.T @ J
        det = float(np.linalg.det(JtJ))
        if det > 0 and math.isfinite(det):
            cov = np.linalg.inv(JtJ) * (rss / dof)
            res.plateau_se = float(math.sqrt(max(cov[0, 0], 0.0)))
            res.k_se = float(math.sqrt(max(cov[1, 1], 0.0)))
            res.cov_kp = float(cov[0, 1])

    res.k = float(k)
    res.plateau = float(plateau)
    res.model = "saturation"
    res.rate = float(k * plateau)
    if math.isfinite(res.k_se) and math.isfinite(res.plateau_se):
        # k_app and Plateau are used as independent parameters (no covariance term)
        var = plateau ** 2 * res.k_se ** 2 + k ** 2 * res.plateau_se ** 2
        res.rate_se = math.sqrt(var) if var > 0 else float("nan")
    res.ok = True
    if math.isfinite(res.r2) and res.r2 < 0.5:
        res.warn = f"poor fit (R^2 = {res.r2:.2f})"
    return res


def last_time_mean(t: np.ndarray, y: np.ndarray) -> float:
    """Mean of y at the largest time point (the normalization reference)."""
    t = np.asarray(t, dtype=float)
    y = np.asarray(y, dtype=float)
    if t.size == 0:
        return float("nan")
    sel = np.isclose(t, float(t.max()))
    vals = y[sel]
    vals = vals[np.isfinite(vals)]
    return float(vals.mean()) if vals.size else float("nan")


def normalized_series(met: "Metabolite", dataset: "Dataset") -> Tuple[np.ndarray, float]:
    """Labeling fraction of one metabolite divided by its last-time-point mean."""
    frac = met.series(dataset.n_cols)
    ref = last_time_mean(dataset.times, frac)
    if not math.isfinite(ref) or ref <= 0:
        return np.full_like(frac, np.nan), ref
    return frac / ref, ref


def fit_all(dataset: Dataset) -> List[FitResult]:
    """Fit L = 1 - exp(-k_app t) for every metabolite (normalized to the last point)."""
    times = dataset.times
    fits: List[FitResult] = []
    for met in dataset.metabolites:
        y, ref = normalized_series(met, dataset)
        fr = fit_kapp(times, y)
        fr.name = met.name
        fr.group = met.group
        fr.label = met.peak_label
        fr.norm = ref
        fr.unit = "relative labeling"
        if not fr.ok and (not math.isfinite(ref) or ref <= 0):
            fr.message = (
                "the labeling fraction at the last time point is "
                f"{ref:.3g}; it must be > 0 to normalize the curve"
            )
        if fr.ok:
            fr.y = y
        fits.append(fr)
    return fits


# ---------------------------------------------------------------------------
# 3. Criterion 1 and flux, on the user-selected pseudo tracer
# ---------------------------------------------------------------------------


def check_criterion1(
    fits: Sequence[FitResult], pt_name: str, members: Sequence[str]
) -> Criterion1Result:
    """k_app(pseudo tracer) >= 4 x k_app of every metabolite used for the flux."""
    out = Criterion1Result(evaluated=True, passed=True)
    by_name = {f.name: f for f in fits}

    pt = by_name.get(pt_name)
    if pt is None:
        out.passed = False
        out.failures.append(f"The pseudo tracer '{pt_name}' was not fitted.")
        return out
    if not pt.ok:
        out.passed = False
        out.failures.append(
            f"The pseudo tracer {pt_name} could not be fitted: {pt.message}"
        )
        return out
    if not members:
        out.passed = False
        out.failures.append("No metabolite was selected for the flux calculation.")
        return out

    pt_k = "instantaneous" if pt.instantaneous else f"{num(pt.k)} h^-1"
    for name in members:
        m = by_name.get(name)
        if m is None:
            out.passed = False
            out.failures.append(f"'{name}' was not fitted, so Criterion 1 cannot be "
                                "evaluated for it.")
            continue
        if not m.ok:
            out.passed = False
            out.failures.append(
                f"The metabolite {name} could not be fitted: {m.message}"
            )
            continue

        if pt.instantaneous:
            out.lines.append(
                f"k_app({pt_name}, pseudo tracer) is instantaneous, so it is >= "
                f"{CRITERION1_FACTOR:g} x k_app({name}) = "
                f"{'instantaneous' if m.instantaneous else num(m.k) + ' h^-1'}  "
                "->  PASS"
            )
            continue

        if m.instantaneous:
            out.passed = False
            out.failures.append(
                f"{name} is labelled instantaneously (k_app is larger than the "
                f"sampling can resolve), so the pseudo tracer {pt_name} "
                f"(k_app = {num(pt.k)} h^-1) cannot be its upstream precursor."
            )
            out.lines.append(
                f"k_app({pt_name}) = {num(pt.k)} h^-1  vs  k_app({name}) = "
                "instantaneous  ->  FAIL"
            )
            continue

        threshold = CRITERION1_FACTOR * m.k
        ok = pt.k >= threshold
        ratio = pt.k / m.k if m.k > 0 else float("inf")
        out.lines.append(
            f"k_app({pt_name}, pseudo tracer) = {pt_k}  vs  "
            f"{CRITERION1_FACTOR:g} x k_app({name}) = "
            f"{CRITERION1_FACTOR:g} x {num(m.k)} = {num(threshold)} h^-1  "
            f"(ratio {num(ratio, 3)})  ->  {'PASS' if ok else 'FAIL'}"
        )
        if not ok:
            out.passed = False
            out.failures.append(
                f"k_app of the pseudo tracer {pt_name} is {num(pt.k)} h^-1, which is "
                f"less than {CRITERION1_FACTOR:g} x k_app of {name} "
                f"({num(m.k)} h^-1); the threshold is {num(threshold)} h^-1."
            )
    return out


def pooled_series(dataset: Dataset, names: Sequence[str]) -> np.ndarray:
    """Sum of the labeled concentrations of the selected metabolites per sample."""
    total = np.zeros(dataset.n_cols, dtype=float)
    for name in names:
        met = dataset.metabolite(name)
        if met is None or met.poolsize is None:
            continue
        total = total + met.concentration(dataset.n_cols)
    return total


def missing_pools(dataset: Dataset, names: Sequence[str]) -> List[str]:
    """Selected metabolites that have no pool size (cannot give a concentration)."""
    out = []
    for name in names:
        met = dataset.metabolite(name)
        if met is None or met.poolsize is None:
            out.append(name)
    return out


def compute_flux(dataset: Dataset, pt_name: str, members: Sequence[str]) -> FluxResult:
    """flux = k_app * Plateau of the summed labeled concentration (single exponential)."""
    if not members:
        raise PipelineError("No metabolite was selected for the flux calculation.")
    missing = missing_pools(dataset, members)
    if missing:
        raise PipelineError(
            "These selected metabolites have no pool size in sheet 2, so a labeled "
            "concentration cannot be computed: " + ", ".join(missing) + "."
        )

    total = pooled_series(dataset, members)
    if not np.any(np.isfinite(total)):
        raise PipelineError("The pooled labeling series contains no usable values.")

    fr = fit_exponential(dataset.times, total)
    if not fr.ok:
        raise PipelineError(
            "The pooled labeled concentration could not be fitted with "
            "L = Plateau x (1 - exp(-k_app x t)): " + fr.message
        )

    flux = fr.k * fr.plateau
    # k_app and Plateau are treated as independent parameters (by design of the
    # method), so the flux variance is propagated WITHOUT the covariance term:
    #     SE(flux)^2 = Plateau^2 * SE(k_app)^2 + k_app^2 * SE(Plateau)^2
    # The fitted covariance is still available as fr.cov_kp / flux.cov_kp.
    var = fr.plateau ** 2 * fr.k_se ** 2 + fr.k ** 2 * fr.plateau_se ** 2
    flux_se = math.sqrt(var) if math.isfinite(var) and var > 0 else float("nan")


    return FluxResult(
        pt=pt_name,
        members=list(members),
        unit=dataset.pool_unit,
        flux=flux,
        flux_se=flux_se,
        k=fr.k,
        k_se=fr.k_se,
        plateau=fr.plateau,
        plateau_se=fr.plateau_se,
        cov_kp=fr.cov_kp,
        r2=fr.r2,
        n=fr.n,
    )


# ---------------------------------------------------------------------------
# 4. Figure
# ---------------------------------------------------------------------------


_PALETTE = [
    "#1f77b4", "#d62728", "#2ca02c", "#ff7f0e", "#9467bd",
    "#8c564b", "#e377c2", "#17becf", "#bcbd22", "#7f7f7f",
]
_MARKERS = ["o", "s", "^", "D", "v", "P", "X", "*", "h", "<"]


def _fit_line(fr: FitResult, t: np.ndarray) -> np.ndarray:
    """Fitted curve: 1 - exp(-k t) for the per-metabolite fits."""
    if fr.model == "saturation":
        return fr.plateau * (-np.expm1(-fr.k * t))
    return -np.expm1(-fr.k * t)


def _fit_legend(fr: FitResult) -> str:
    if not fr.ok:
        return fr.name
    if fr.instantaneous:
        return f"{fr.name}: k_app $>$ limit (instantaneous)"
    return f"{fr.name}: k_app = {num(fr.k, 3)} h$^{{-1}}$"


def build_figure(dataset: Dataset, fits: Sequence[FitResult],
                 figsize: Tuple[float, float] = (9.4, 4.4),
                 dpi: int = 100) -> "Figure":
    """Normalized labeling curves: one measured series and one fitted curve each."""
    if Figure is None:
        raise RuntimeError("matplotlib is required for plotting: " + str(_MPL_ERR))

    fig = Figure(figsize=figsize, dpi=dpi)
    try:
        fig.set_layout_engine("tight")
    except Exception:  # matplotlib < 3.6
        pass
    ax = fig.add_subplot(111)

    by_name = {f.name: f for f in fits}
    t_end = max(dataset.unique_times()) if dataset.columns else 1.0
    t_grid = np.linspace(0.0, t_end * 1.03, 400)
    times = dataset.times

    for i, met in enumerate(dataset.metabolites):
        fr = by_name.get(met.name)
        y, ref = normalized_series(met, dataset)
        if not math.isfinite(ref) or ref <= 0:
            continue
        color = _PALETTE[i % len(_PALETTE)]
        marker = _MARKERS[i % len(_MARKERS)]
        ut, mean, sd = mean_sd_by_time(times, y)
        # observed data: symbols + SD bars only, never connected by a line
        ax.errorbar(ut, mean, yerr=sd, fmt=marker, ms=5.5, color=color, mfc="white",
                    mew=1.3, capsize=3, elinewidth=1.0, linestyle="none", zorder=3)
        if fr is not None and fr.ok:
            ax.plot(t_grid, _fit_line(fr, t_grid), color=color, lw=1.6, zorder=2,
                    label=_fit_legend(fr))
        else:
            ax.plot([], [], marker=marker, color=color, mfc="white", mew=1.3,
                    linestyle="none", label=f"{met.name} (not fitted)")

    ax.set_xlabel(
        "Time after tracer addition (h)\n"
        "(symbols: mean \u00b1 SD of the replicates;  lines: least-squares fit)"
    )
    ax.set_ylabel("Normalized labeling\n(L / mean of the last time point)", fontsize=9.5)
    ax.set_title("SP-KFP fits   L = 1 - exp(-k_app x t)", fontsize=11.5)
    ax.grid(True, ls=":", lw=0.6, alpha=0.7)
    ax.set_xlim(0, t_end * 1.03)
    ax.set_ylim(bottom=0)
    ax.axhline(1.0, color="#999999", lw=0.8, ls="--", zorder=1)
    ax.legend(loc="lower right", fontsize=7.4, framealpha=0.92, borderpad=0.4,
              handlelength=2.2)
    return fig


# ---------------------------------------------------------------------------
# 5. Export
# ---------------------------------------------------------------------------


def export_results(path: str, dataset: Dataset, fits: Sequence[FitResult],
                   crit1: Optional[Criterion1Result],
                   flux: Optional[FluxResult]) -> None:
    """Write the fitted k_app values, the criterion-1 report and the flux."""
    if openpyxl is None:
        raise RuntimeError("openpyxl is required to export results")
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Fits"
    ws.append([
        "Metabolite", "Labeled form", "k_app (h^-1)", "SE k_app (h^-1)",
        "R^2", "n", "Normalization (mean of the last time point)",
        "Pool size", "Note",
    ])
    for c in ws[1]:
        c.font = Font(bold=True)

    for met in dataset.metabolites:
        fr = next((f for f in fits if f.name == met.name), None)
        if fr is None:
            continue
        note = fr.message if not fr.ok else (fr.warn or "")
        ws.append([
            fr.name, fr.label,
            fr.k if fr.ok else None, fr.k_se if fr.ok else None,
            fr.r2, fr.n, fr.norm, met.poolsize, note,
        ])

    for i, width in enumerate([18, 26, 13, 14, 8, 6, 34, 12, 60], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws2 = wb.create_sheet("Criterion 1")
    ws2.append(["Criterion 1 (precursor sufficiency): k_app(pseudo tracer) "
                ">= 4 x k_app(every metabolite used for the flux)"])
    if crit1 is None:
        ws2.append(["Result", "not evaluated"])
    else:
        ws2.append(["Result", "PASS" if crit1.passed else "FAIL"])
        for line in crit1.lines:
            ws2.append([line])
        for line in crit1.failures:
            ws2.append(["FAILURE: " + line])
    ws2.column_dimensions["A"].width = 120

    ws3 = wb.create_sheet("Flux")
    ws3.append(["SP-KFP flux"])
    if flux is not None:
        ws3.append(["Pseudo tracer", flux.pt])
        ws3.append(["Metabolites summed", ", ".join(flux.members)])
        ws3.append([])
        ws3.append(["Flux", flux.flux, "SE", flux.flux_se, f"{flux.unit}/h"])
        ws3.append(["k_app (h^-1)", flux.k, "SE", flux.k_se])
        ws3.append(["Plateau", flux.plateau, "SE", flux.plateau_se, flux.unit])
        ws3.append(["R^2", flux.r2])
        ws3.append(["n data points", flux.n])
        if flux.note:
            ws3.append(["Note", flux.note])
    else:
        ws3.append(["Flux", "not calculated"])
    ws3.append([])
    for note in dataset.notes:
        ws3.append(["Note", note])
    ws3.column_dimensions["A"].width = 30
    ws3.column_dimensions["B"].width = 60

    wb.save(path)


# ---------------------------------------------------------------------------
# 6. GUI
# ---------------------------------------------------------------------------


def _fmt_num(x, digits=4, na="n/a"):
    try:
        v = float(x)
    except (TypeError, ValueError):
        return na
    if not math.isfinite(v):
        return na
    return f"{v:.{digits}g}"


class SpkfpApp:
    """Tkinter front end: submit -> criterion 1 -> criteria 2/3 -> flux -> output."""

    def __init__(self, root):
        self.root = root
        self.dataset: Optional[Dataset] = None
        self.fits: List[FitResult] = []
        self.crit1: Optional[Criterion1Result] = None
        self.flux: Optional[FluxResult] = None
        self.figure = None
        self.canvas = None

        root.title("SP-KFP flux calculator  -  non-stationary isotope tracing")
        root.geometry("1360x1120")
        root.minsize(1100, 780)
        root.columnconfigure(0, weight=1)
        root.rowconfigure(3, weight=1, minsize=260)   # figure grows

        self._build_toolbar()      # row 0
        self._build_hint()         # row 1
        self._build_table()        # row 2
        self._build_figure_area()  # row 3 (expands)
        self._build_action_bar()   # row 4
        self._build_output_area()  # row 5
        self._build_statusbar()    # row 6
        self._set_state("idle")

    # ---------------- widgets ----------------
    def _build_toolbar(self):
        bar = ttk.Frame(self.root, padding=(12, 10, 12, 2))
        bar.grid(row=0, column=0, sticky="ew")
        bar.columnconfigure(2, weight=1)

        self.btn_load = ttk.Button(bar, text="Load data (.xlsx) \u2026",
                                   command=self.on_load)
        self.btn_load.grid(row=0, column=0, padx=(0, 8))

        self.btn_export = ttk.Button(bar, text="Export results \u2026",
                                     command=self.on_export, state="disabled")
        self.btn_export.grid(row=0, column=1, padx=(0, 8))

        self.lbl_info = ttk.Label(bar, text="No data loaded.", anchor="e")
        self.lbl_info.grid(row=0, column=2, sticky="e")

    def _build_hint(self):
        ttk.Label(
            self.root,
            text=("Sheet 1: metabolite, isotopeLabel, formula, then the sample columns "
                  "'0_1', '0.5_2', ... (time in hours).     "
                  "Sheet 2: unit in B1, then metabolite and pool size."),
            padding=(12, 0, 12, 6),
            foreground="#555555",
        ).grid(row=1, column=0, sticky="ew")

    def _build_table(self):
        frame = ttk.LabelFrame(
            self.root,
            text="  Fitted k_app per metabolite (L = 1 - exp(-k_app x t), "
                 "normalized to the last time point)  ",
            padding=6)
        frame.grid(row=2, column=0, sticky="ew", padx=12, pady=(2, 6))
        frame.columnconfigure(0, weight=1)

        cols = ("name", "form", "k", "kse", "r2", "n", "norm", "pool", "note")
        heads = ("Metabolite", "Labeled form", "k_app (h\u207b\u00b9)", "SE(k_app)",
                 "R\u00b2", "n", "L_last (raw)", "Pool size", "Note")
        widths = (150, 150, 110, 95, 55, 40, 110, 105, 330)

        self.tree = ttk.Treeview(frame, columns=cols, show="headings", height=6)
        for c, h, w in zip(cols, heads, widths):
            self.tree.heading(c, text=h)
            self.tree.column(
                c, width=w, minwidth=45, stretch=(c == "note"),
                anchor="w" if c in ("name", "form", "note") else "e")
        self.tree.grid(row=0, column=0, sticky="ew")
        sb = ttk.Scrollbar(frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(xscrollcommand=sb.set)
        sb.grid(row=1, column=0, sticky="ew")

    def _build_figure_area(self):
        frame = ttk.LabelFrame(self.root, text="  Fitted labeling curves  ", padding=6)
        frame.grid(row=3, column=0, sticky="nsew", padx=12, pady=(0, 6))
        frame.columnconfigure(0, weight=1)
        frame.rowconfigure(0, weight=1)
        self.fig_frame = frame

        self.fig_holder = ttk.Frame(frame)
        self.fig_holder.grid(row=0, column=0, sticky="nsew")
        self.fig_holder.columnconfigure(0, weight=1)
        self.fig_holder.rowconfigure(0, weight=1)
        self.fig_placeholder = ttk.Label(
            self.fig_holder,
            text="\n\nNo figure yet.\n\nLoad a data file to audit the workbook, check "
                 "Criterion 1 and draw the fitted curves.",
            anchor="center", justify="center", foreground="#666666",
        )
        self.fig_placeholder.grid(row=0, column=0, sticky="nsew")

        row = ttk.Frame(frame)
        row.grid(row=1, column=0, sticky="ew", pady=(4, 0))
        self.btn_save = ttk.Button(row, text="Save figure as image \u2026",
                                   command=self.on_save_figure, state="disabled")
        self.btn_save.pack(side="left")
        self.lbl_c1 = ttk.Label(
            row, text="Criterion 1 is checked when you choose the pseudo tracer.",
            foreground="#444444")
        self.lbl_c1.pack(side="left", padx=14)

    def _build_action_bar(self):
        bar = ttk.Frame(self.root, padding=(12, 0, 12, 8))
        bar.grid(row=4, column=0, sticky="ew")
        self.btn_flux = ttk.Button(bar, text="Calculate flux",
                                   command=self.on_calculate_flux, state="disabled")
        self.btn_flux.pack(side="left")
        ttk.Label(bar,
                  text="(opens a dialog where you pick the pseudo tracer and the "
                       "metabolites to pool, then fills the flux panel below)",
                  foreground="#555555").pack(side="left", padx=12)

    def _build_output_area(self):
        frame = ttk.LabelFrame(self.root, text="  Flux result  ", padding=10)
        frame.grid(row=5, column=0, sticky="ew", padx=12, pady=(0, 6))
        frame.columnconfigure(0, weight=1)
        self.out_frame = frame
        self.out_body = ttk.Frame(frame)
        self.out_body.grid(row=0, column=0, sticky="ew")
        self.out_body.columnconfigure(0, weight=1)
        self._set_output_text(
            "The flux will be shown here after you click \u201cCalculate flux\u201d.", "")

    def _build_statusbar(self):
        self.status = ttk.Label(self.root, text="Ready.", relief="sunken", anchor="w",
                                padding=(8, 3))
        self.status.grid(row=6, column=0, sticky="ew")

    def _set_output_text(self, main: str, detail: str):
        for w in self.out_body.winfo_children():
            w.destroy()
        self.out_flux_lbl = ttk.Label(self.out_body, text=main, font=("Segoe UI", 13, "bold"),
                                      foreground="#0b3d6b", wraplength=1300)
        self.out_flux_lbl.grid(row=0, column=0, sticky="w")
        self.out_detail = ttk.Label(self.out_body, text=detail, justify="left",
                                    foreground="#333333", wraplength=1240,
                                    font=("Segoe UI", 9))
        self.out_detail.grid(row=1, column=0, sticky="w", pady=(6, 0))

    def _reset_view(self):
        """Blank the table, the figure and the flux panel (back to the idle state)."""
        for item in self.tree.get_children():
            self.tree.delete(item)
        for w in self.fig_holder.winfo_children():
            w.destroy()
        self.figure = None
        self.canvas = None
        self.fits = []
        self.crit1 = None
        self.flux = None
        self.fig_placeholder = ttk.Label(
            self.fig_holder,
            text="\n\nNo figure yet.\n\nLoad a data file to fit one k_app per "
                 "metabolite\nand draw the normalized labeling curves.",
            anchor="center", justify="center", foreground="#666666",
        )
        self.fig_placeholder.grid(row=0, column=0, sticky="nsew")
        self.lbl_c1.configure(
            text="Criterion 1 is checked when you choose the pseudo tracer.",
            foreground="#444444")
        self._set_output_text(
            "The flux will be shown here after you click \u201cCalculate flux\u201d.", "")
        self._set_state("idle")

    def _set_state(self, state: str):
        loaded = state in ("loaded", "flux")
        self.btn_save.configure(state="normal" if loaded else "disabled")
        self.btn_export.configure(state="normal" if state == "flux" else "disabled")
        self.btn_flux.configure(state="normal" if loaded else "disabled")

    def _status(self, text: str):
        self.status.configure(text=text)

    # ---------------- dialogs ----------------
    def _scroll_dialog(self, title: str, intro: str, lines: Sequence[str],
                       buttons: Sequence[Tuple[str, object]], width: int = 96,
                       height: int = 20):
        """Modal, scrollable message box.  Returns the value of the clicked button."""
        win = tk.Toplevel(self.root)
        win.title(title)
        win.transient(self.root)
        win.resizable(True, True)
        win.columnconfigure(0, weight=1)
        win.rowconfigure(1, weight=1)

        head = ttk.Label(win, text=intro, padding=(12, 10, 12, 6), justify="left",
                         wraplength=760)
        head.grid(row=0, column=0, sticky="ew")

        box = ttk.Frame(win, padding=(12, 0, 12, 0))
        box.grid(row=1, column=0, sticky="nsew")
        box.columnconfigure(0, weight=1)
        box.rowconfigure(0, weight=1)
        txt = tk.Text(box, wrap="word", width=width, height=height, relief="solid",
                      borderwidth=1, padx=8, pady=6, background="#fbfbfb")
        txt.grid(row=0, column=0, sticky="nsew")
        sb = ttk.Scrollbar(box, orient="vertical", command=txt.yview)
        sb.grid(row=0, column=1, sticky="ns")
        txt.configure(yscrollcommand=sb.set)
        for line in lines:
            txt.insert("end", line + "\n")
        txt.configure(state="disabled")

        bar = ttk.Frame(win, padding=(12, 8, 12, 12))
        bar.grid(row=2, column=0, sticky="ew")
        result = {"value": buttons[0][1] if buttons else None}

        def close(val):
            result["value"] = val
            win.destroy()

        for text, value in reversed(list(buttons)):
            ttk.Button(bar, text=text, command=lambda v=value: close(v)).pack(
                side="right", padx=(8, 0))
        win.protocol("WM_DELETE_WINDOW", lambda: close(buttons[-1][1] if buttons else None))

        win.update_idletasks()
        try:
            x = self.root.winfo_rootx() + max(0, (self.root.winfo_width() - win.winfo_width()) // 2)
            y = self.root.winfo_rooty() + max(0, (self.root.winfo_height() - win.winfo_height()) // 3)
            win.geometry(f"+{x}+{y}")
        except Exception:
            pass
        win.grab_set()
        win.wait_window()
        return result["value"]

    def _error(self, title: str, intro: str, lines: Sequence[str]):
        self._scroll_dialog(title, intro, lines, [("OK", True)])

    # ---------------- step 1: load ----------------
    def on_load(self):
        initial = os.path.dirname(self.dataset.path) if self.dataset else os.path.expanduser("~")
        path = filedialog.askopenfilename(
            title="Select the SP-KFP data workbook",
            initialdir=initial,
            filetypes=[("Excel workbook", "*.xlsx"), ("All files", "*.*")],
        )
        if not path:
            return

        while True:
            try:
                dataset = load_dataset(path)
                break
            except ValidationError as exc:
                self._error(
                    exc.title,
                    "The submitted workbook does not meet the SP-KFP data "
                    "requirements. The file was not accepted.",
                    exc.problems + ["", "Please correct the file and submit it again."],
                )
                self._status("File rejected: " + exc.problems[0])
                path = filedialog.askopenfilename(
                    title="Please submit the corrected data workbook",
                    initialdir=os.path.dirname(path),
                    filetypes=[("Excel workbook", "*.xlsx"), ("All files", "*.*")],
                )
                if not path:
                    self._status("Ready. No data loaded.")
                    return
            except Exception as exc:  # unexpected problem
                self._error("Data file rejected",
                            "The workbook could not be processed.",
                            [str(exc), "", "The file was not accepted. Please submit "
                                           "another file."])
                self._status("File rejected: " + str(exc))
                path = filedialog.askopenfilename(
                    title="Please submit the data workbook",
                    initialdir=os.path.dirname(path),
                    filetypes=[("Excel workbook", "*.xlsx"), ("All files", "*.*")],
                )
                if not path:
                    self._status("Ready. No data loaded.")
                    return

        self.dataset = dataset
        self._reset_view()
        self.lbl_info.configure(text=dataset.summary())

        if dataset.notes:
            self._scroll_dialog(
                "Data accepted - with notes",
                "The workbook passed the audit. Please review the following notes:",
                dataset.notes + ["", "Click OK to continue."],
                [("OK", True)],
            )

        self._run_analysis()

    def _run_analysis(self):
        assert self.dataset is not None
        dataset = self.dataset
        self._status("Fitting one k_app per metabolite ...")
        self.root.update_idletasks()

        fits = fit_all(dataset)
        self.fits = fits

        failed = [f for f in fits if not f.ok]
        self._populate_table()
        self._draw_figure()

        if failed:
            lines = [f"  {f.name}: {f.message}" for f in failed]
            self._error(
                "Some metabolites could not be fitted",
                "A curve L = 1 - exp(-k_app x t) could not be fitted to the following "
                "metabolite(s). They are shown as \"not fitted\" and cannot be used "
                "in the flux calculation.",
                lines + ["", "Please check their labeling values (at least 3 time points "
                              "and a labeling fraction > 0 at the last time point are "
                              "required)."],
            )
            self._set_state("loaded")
            self.lbl_c1.configure(
                text=f"{len(failed)} metabolite(s) could not be fitted.",
                foreground="#a11212")
            self._status("Fitted with warnings: " + ", ".join(f.name for f in failed))
            return

        self._set_state("loaded")
        self.lbl_c1.configure(
            text="Criterion 1 is checked when you choose the pseudo tracer.",
            foreground="#444444")
        self._status(
            "k_app fitted for all metabolites. Click \u201cCalculate flux\u201d to choose "
            "the pseudo tracer and the metabolites to pool.")

    # ---------------- table + figure ----------------
    def _populate_table(self):
        assert self.dataset is not None
        for item in self.tree.get_children():
            self.tree.delete(item)
        for met in self.dataset.metabolites:
            fr = next((f for f in self.fits if f.name == met.name), None)
            if fr is None:
                continue
            if fr.ok:
                if fr.instantaneous:
                    k_txt, kse_txt = "n.i. (inst.)", "-"
                else:
                    k_txt, kse_txt = _fmt_num(fr.k), _fmt_num(fr.k_se)
                note = fr.warn or ""
            else:
                k_txt = kse_txt = "n/a"
                note = fr.message
            pool = _fmt_num(met.poolsize) if met.poolsize is not None else "not given"
            self.tree.insert("", "end", values=(
                met.name, fr.label, k_txt, kse_txt, _fmt_num(fr.r2, 3), fr.n,
                _fmt_num(fr.norm, 4), pool, note,
            ))

    def _draw_figure(self):
        assert self.dataset is not None
        try:
            from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
        except Exception as exc:
            self._error("Plotting unavailable",
                        "matplotlib could not be loaded, so the fitted curves cannot "
                        "be drawn.",
                        [str(exc)])
            return

        for w in self.fig_holder.winfo_children():
            w.destroy()
        self.figure = build_figure(self.dataset, self.fits)
        self.canvas = FigureCanvasTkAgg(self.figure, master=self.fig_holder)
        widget = self.canvas.get_tk_widget()
        widget.grid(row=0, column=0, sticky="nsew")
        self.canvas.draw()
        self.fig_holder.bind("<Configure>", self._sync_figure_size)
        self.root.after(60, self._sync_figure_size)

    def _sync_figure_size(self, event=None):
        """Keep the matplotlib figure exactly as large as its Tk cell (HiDPI safe)."""
        if self.figure is None or self.canvas is None:
            return
        w, h = self.fig_holder.winfo_width(), self.fig_holder.winfo_height()
        if w < 120 or h < 120:
            return
        dpi = float(self.figure.dpi) or 100.0
        cur_w, cur_h = self.figure.get_size_inches()
        if abs(cur_w * dpi - w) <= 2 and abs(cur_h * dpi - h) <= 2:
            return
        self.figure.set_size_inches(w / dpi, h / dpi, forward=False)
        self.canvas.draw_idle()

    # ---------------- save / export ----------------
    def on_save_figure(self):
        if self.figure is None:
            return
        default = "SPKFP_fitting.png"
        if self.dataset:
            stem = os.path.splitext(os.path.basename(self.dataset.path))[0]
            default = f"{stem}_SPKFP_fits.png"
        path = filedialog.asksaveasfilename(
            title="Save the fitted curves",
            defaultextension=".png",
            initialfile=default,
            filetypes=[("PNG image", "*.png"), ("TIFF image", "*.tif"),
                       ("PDF document", "*.pdf"), ("SVG image", "*.svg")],
        )
        if not path:
            return
        try:
            # re-render at publication aspect ratio instead of the (wide) screen
            # canvas, then write it at high resolution
            fig = build_figure(self.dataset, self.fits, figsize=(9.8, 5.3), dpi=300)
            fig.savefig(path, dpi=300, bbox_inches="tight")
            self._status(f"Figure saved to {path}")
        except Exception as exc:
            self._error("Could not save the figure", "Saving the image failed.", [str(exc)])

    def on_export(self):
        if self.dataset is None or self.crit1 is None:
            return
        path = filedialog.asksaveasfilename(
            title="Export the results",
            defaultextension=".xlsx",
            initialfile="SPKFP_results.xlsx",
            filetypes=[("Excel workbook", "*.xlsx")],
        )
        if not path:
            return
        try:
            export_results(path, self.dataset, self.fits, self.crit1, self.flux)
            self._status(f"Results exported to {path}")
        except Exception as exc:
            self._error("Could not export the results", "Exporting failed.", [str(exc)])

    # ---------------- step 3/4: criteria 2-3 and flux ----------------
    def _select_flux_inputs(self) -> Optional[Tuple[str, List[str]]]:
        """Dialog: pick ONE pseudo tracer and the metabolites to pool for the flux."""
        assert self.dataset is not None
        usable = [m for m in self.dataset.metabolites
                  if any(f.name == m.name and f.ok for f in self.fits)]
        if not usable:
            self._error("Nothing to choose",
                        "No metabolite could be fitted, so the flux cannot be "
                        "calculated.", [])
            return None

        win = tk.Toplevel(self.root)
        win.title("Choose the pseudo tracer and the metabolites to pool")
        win.transient(self.root)
        win.resizable(True, False)

        ttk.Label(
            win,
            text="Pseudo tracer (select exactly one) - the metabolite X whose k_app is "
                 "compared with every metabolite used for the flux:",
            padding=(14, 12, 14, 4), wraplength=700, justify="left",
            font=("Segoe UI", 10, "bold"),
        ).grid(row=0, column=0, sticky="ew")

        top = ttk.Frame(win, padding=(14, 0, 14, 6))
        top.grid(row=1, column=0, sticky="ew")
        pt_var = tk.StringVar(value="")
        pt_box = ttk.Combobox(top, textvariable=pt_var, state="readonly", width=40,
                              values=[f"{f.name}   (k_app = {num(f.k)} h\u207b\u00b9)"
                                      if not f.instantaneous else
                                      f"{f.name}   (k_app > limit, instantaneous)"
                                      for f in self.fits if f.ok])
        pt_box.pack(side="left")
        ttk.Label(top, text="    k_app of the pseudo tracer must be \u2265 4 \u00d7 k_app "
                            "of every metabolite used for the flux.",
                  foreground="#555555").pack(side="left")

        mid = ttk.LabelFrame(
            win, padding=8,
            text="  Metabolites to pool for the flux  (one or more; they must all be "
                 "DOWNSTREAM of the pseudo tracer)  ")
        mid.grid(row=2, column=0, sticky="ew", padx=14, pady=4)
        mid.columnconfigure(0, weight=1)

        ttk.Label(
            mid,
            text="Check every metabolite whose labeled concentration should be summed. "
                 "Criteria 2 and 3 of the notes apply to this group: they are merged into "
                 "one pool only if they rapidly equilibrate / are the major downstream "
                 "branch.",
            wraplength=700, justify="left", foreground="#7a3b00",
        ).grid(row=0, column=0, sticky="w", pady=(0, 6))

        listbox_frame = ttk.Frame(mid)
        listbox_frame.grid(row=1, column=0, sticky="ew")
        listbox_frame.columnconfigure(0, weight=1)
        canvas = tk.Canvas(listbox_frame, height=190, highlightthickness=0)
        vsb = ttk.Scrollbar(listbox_frame, orient="vertical", command=canvas.yview)
        canvas.configure(yscrollcommand=vsb.set)
        canvas.grid(row=0, column=0, sticky="ew")
        vsb.grid(row=0, column=1, sticky="ns")
        inner = ttk.Frame(canvas)
        canvas.create_window((0, 0), window=inner, anchor="nw")
        inner.bind("<Configure>",
                   lambda e: canvas.configure(scrollregion=canvas.bbox("all")))

        checks: List[Tuple[str, tk.BooleanVar, ttk.Checkbutton]] = []
        for met in usable:
            fr = next(f for f in self.fits if f.name == met.name)
            has_pool = met.poolsize is not None
            if fr.instantaneous:
                label = f"{met.name}   (k_app > limit, instantaneous)"
            else:
                label = f"{met.name}   (k_app = {num(fr.k)} h\u207b\u00b9)"
            if not has_pool:
                label += "   - no pool size, cannot be used"
            var = tk.BooleanVar(value=False)
            chk = ttk.Checkbutton(inner, variable=var, text=label)
            chk.grid(sticky="w", pady=1)
            if not has_pool:
                chk.configure(state="disabled")
            checks.append((met.name, var, chk))

        def on_pt_change(*_):
            chosen = pt_var.get().split("   (")[0]
            for name, var, chk in checks:
                fr = next((f for f in self.fits if f.name == name), None)
                if name == chosen:
                    var.set(False)
                    chk.configure(state="disabled")
                elif fr is not None and fr.ok and \
                        self.dataset.metabolite(name).poolsize is not None:
                    chk.configure(state="normal")

        pt_box.bind("<<ComboboxSelected>>", on_pt_change)

        warn = ttk.Label(
            win,
            text="Reminder: the metabolites you tick must be DOWNSTREAM of the pseudo "
                 "tracer. The flux is calculated from the sum of their labeled "
                 "concentrations, fitted with L = Plateau x (1 - exp(-k_app x t)), and "
                 "flux = k_app x Plateau.",
            padding=(14, 6, 14, 6), wraplength=700, justify="left",
            foreground="#8a1c1c",
        )
        warn.grid(row=3, column=0, sticky="ew")

        bar = ttk.Frame(win, padding=(14, 6, 14, 12))
        bar.grid(row=4, column=0, sticky="ew")
        result: Dict[str, object] = {"value": None}

        def accept():
            name = pt_var.get().split("   (")[0].strip()
            members = [n for n, v, _c in checks if v.get()]
            if not name or not members:
                return
            result["value"] = (name, members)
            win.destroy()

        def cancel():
            result["value"] = None
            win.destroy()

        ok_btn = ttk.Button(bar, text="Check Criterion 1 and calculate flux",
                            command=accept, state="disabled")
        ok_btn.pack(side="right", padx=(8, 0))
        ttk.Button(bar, text="Cancel", command=cancel).pack(side="right")

        def refresh(*_):
            members = [n for n, v, _c in checks if v.get()]
            ready = bool(pt_var.get().strip()) and bool(members)
            ok_btn.configure(state="normal" if ready else "disabled")

        for _n, var, _c in checks:
            var.trace_add("write", refresh)
        pt_var.trace_add("write", refresh)
        refresh()

        win.protocol("WM_DELETE_WINDOW", cancel)
        win.grab_set()
        win.wait_window()
        return result["value"]  # type: ignore[return-value]

    def on_calculate_flux(self):
        if self.dataset is None or not self.fits:
            return
        selection = self._select_flux_inputs()
        if not selection:
            self._status("Flux calculation cancelled.")
            return
        pt_name, members = selection

        missing = missing_pools(self.dataset, members)
        if missing:
            self._error(
                "Pool size missing",
                "These metabolites have no pool size in sheet 2, so their labeled "
                "concentration cannot be computed:",
                [", ".join(missing), "",
                 "Please add their pool sizes to the second worksheet and load the file "
                 "again, or exclude them from the flux."],
            )
            self._status("Flux calculation aborted: missing pool size.")
            return

        crit1 = check_criterion1(self.fits, pt_name, members)
        self.crit1 = crit1
        if not crit1.passed:
            self._error(
                "Criterion 1 is not satisfied",
                "Criterion 1 (precursor sufficiency) requires that the k_app of the "
                f"pseudo tracer ({pt_name}) is at least 4 x the k_app of every "
                "metabolite used for the flux. This is not met, so the "
                "single-exponential approximation would carry an error larger than "
                "25 % and the flux is not calculated.",
                crit1.failures + ["", "Comparison details:"] + crit1.lines,
            )
            self.lbl_c1.configure(text="Criterion 1: FAILED", foreground="#a11212")
            self._status("Criterion 1 not satisfied - flux was not calculated.")
            return

        self.lbl_c1.configure(text="Criterion 1: PASS", foreground="#146c2e")
        self._status("Calculating the pooled flux ...")
        self.root.update_idletasks()
        try:
            flux = compute_flux(self.dataset, pt_name, members)
        except PipelineError as exc:
            self._error("Flux could not be calculated", "The pooled fit failed.",
                        [str(exc)])
            self._status("Flux calculation failed.")
            return
        except Exception as exc:
            self._error("Flux could not be calculated", "An unexpected error occurred.",
                        [repr(exc)])
            self._status("Flux calculation failed.")
            return

        self.flux = flux
        unit = f"{flux.unit}/h"
        main = (f"Flux = {_fmt_num(flux.flux)} \u00b1 {_fmt_num(flux.flux_se)} {unit}"
                f"    (k_app \u00d7 Plateau of the summed labeled concentration)")
        detail = "\n".join([
            f"Pseudo tracer: {flux.pt}   |   summed metabolites ({len(flux.members)}): "
            f"{', '.join(flux.members)}",
            f"   k_app (pooled) = {_fmt_num(flux.k)} \u00b1 {_fmt_num(flux.k_se)} "
            f"h\u207b\u00b9     Plateau = {_fmt_num(flux.plateau)} \u00b1 "
            f"{_fmt_num(flux.plateau_se)} {flux.unit}",
            f"   R\u00b2 = {_fmt_num(flux.r2, 3)},  n = {flux.n} sample points "
            "(SE from SE(k_app) and SE(Plateau), treated as independent)",
            "",
            "Criterion 1 (k_app of the pseudo tracer \u2265 4 \u00d7 k_app of every "
            "pooled metabolite):",
        ] + ["   " + line for line in self.crit1.lines] + (
            ["", flux.note] if flux.note else []
        ))
        self._set_output_text(main, detail)
        self._set_state("flux")
        self._status("Done. Flux = " + _fmt_num(flux.flux) + " " + unit)


# ---------------------------------------------------------------------------
# 7. Entry points
# ---------------------------------------------------------------------------


def _check_cli(path: str, pt_name: str = "", flux_names: Optional[Sequence[str]] = None) -> int:
    """Headless run: audit, fit, and (optionally) criterion 1 + flux as ASCII text."""
    flux_names = list(flux_names or [])
    try:
        dataset = load_dataset(path)
    except ValidationError as exc:
        print("REJECTED")
        for p in exc.problems:
            print("  - " + p)
        return 2

    print("ACCEPTED")
    print("  " + dataset.summary())
    for n in dataset.notes:
        print("  note: " + n)

    fits = fit_all(dataset)
    print("")
    print(f"{'metabolite':18} {'form':24} {'k_app':>10} {'SE':>10} {'R2':>6} {'n':>3}")
    print("-" * 76)
    for met in dataset.metabolites:
        fr = next((f for f in fits if f.name == met.name), None)
        if fr is None:
            continue
        if fr.ok:
            k_txt = "n.i.(inst)" if fr.instantaneous else num(fr.k)
            se_txt = "-" if fr.instantaneous else num(fr.k_se)
        else:
            k_txt, se_txt = "n/a", "-"
        print(f"{met.name[:18]:18} {fr.label[:24]:24} {k_txt:>10} {se_txt:>10} "
              f"{num(fr.r2, 3):>6} {fr.n:>3}"
              + (f"  [{fr.message or fr.warn}]" if not fr.ok else ""))

    bad = [f for f in fits if not f.ok]
    if bad:
        print("")
        for f in bad:
            print(f"  FIT FAILED: {f.name}: {f.message}")
        return 3

    if not pt_name or not flux_names:
        print("")
        print("Criterion 1 and the flux need a pseudo tracer and the metabolites to "
              "pool:")
        print("  spkfp_app.py --check <workbook.xlsx> --pt <metabolite> "
              "--flux <metabolite>[,<metabolite>...]")
        return 0

    crit1 = check_criterion1(fits, pt_name, flux_names)
    print("")
    print(f"Pseudo tracer: {pt_name}")
    print(f"Metabolites pooled: {', '.join(flux_names)}")
    print(f"Criterion 1: {'PASS' if crit1.passed else 'FAIL'}")
    for line in crit1.lines:
        print("  " + line.replace("\u00b9", "^-1"))
    for line in crit1.failures:
        print("  FAILURE: " + line)
    if not crit1.passed:
        return 4

    flux = compute_flux(dataset, pt_name, flux_names)
    print("")
    print(f"  k_app   = {num(flux.k)} +/- {num(flux.k_se)} h^-1")
    print(f"  Plateau = {num(flux.plateau)} +/- {num(flux.plateau_se)} {flux.unit}")
    print(f"  R2      = {num(flux.r2, 3)}   n = {flux.n}")
    print(f"  FLUX    = {num(flux.flux)} +/- {num(flux.flux_se)} {flux.unit}/h")
    if flux.note:
        print("  note: " + flux.note)
    return 0


def script_path() -> str:
    """Absolute path of this script (empty string when it cannot be resolved)."""
    path = globals().get("__file__", "")
    try:
        return os.path.abspath(path) if path else "spkfp_app.py"
    except Exception:
        return "spkfp_app.py"


def _native_message(title: str, message: str) -> None:
    """Show a message box without tkinter (used when tkinter itself is missing)."""
    _log(message)
    if os.name == "nt":
        try:
            import ctypes

            # MB_ICONERROR | MB_SETFOREGROUND; keeps the text visible after a
            # double-click, where the console window would close immediately
            ctypes.windll.user32.MessageBoxW(None, message, title, 0x10 | 0x10000)
            return
        except Exception:
            pass
    try:  # macOS / Linux without a display: at least try a terminal dialog
        subprocess.run(["zenity", "--error", "--text", message, "--title", title],
                       timeout=300, check=False, **(_subprocess_flags()))
    except Exception:
        pass


def _report_missing_tkinter() -> int:
    msg = (
        "SP-KFP needs 'tkinter' for its window, but this Python installation does "
        "not contain it.\n\n"
        "tkinter cannot be installed with pip - it belongs to Python itself:\n"
        "  * Windows: run the Python installer again, choose \"Modify\" and tick\n"
        "             \"tcl/tk and IDLE\";\n"
        "  * macOS:   use the installer from python.org (it includes tkinter);\n"
        "  * Linux:   sudo apt install python3-tk   (or your distribution's "
        "equivalent).\n\n"
        "The command-line mode does not need a window, for example:\n"
        f'  "{sys.executable}" "{script_path()}" --check data.xlsx\n\n'
        f"Python in use: {sys.executable}"
    )
    _native_message("SP-KFP - tkinter is missing", msg)
    return 1


def _report_missing_gui_module(name: str, message: str) -> int:
    advice = (
        f"{message}\n\n"
        "It can be installed once with:\n"
        f'  "{sys.executable}" -m pip install {name}\n\n'
        "Automatic installation is disabled only when SPKFP_NO_AUTO_INSTALL is set; "
        "remove that variable and start the program again."
    )
    _native_message(f"SP-KFP - {name} is missing", advice)
    return 1


def main(argv: Optional[Sequence[str]] = None) -> int:
    argv = list(sys.argv[1:] if argv is None else argv)

    if argv and argv[0] in ("--check", "-c"):
        if len(argv) < 2:
            print("usage: spkfp_app.py --check <workbook.xlsx> "
                  "[--pt NAME] [--flux NAME[,NAME...]]")
            return 1
        pt_name = ""
        flux_names: List[str] = []
        rest = argv[1:]
        path = rest[0]
        i = 1
        while i < len(rest):
            if rest[i] == "--pt" and i + 1 < len(rest):
                pt_name = rest[i + 1]
                i += 2
            elif rest[i] == "--flux" and i + 1 < len(rest):
                flux_names = [x.strip() for x in rest[i + 1].split(",") if x.strip()]
                i += 2
            else:
                i += 1
        return _check_cli(path, pt_name, flux_names)
    if argv and argv[0] in ("-h", "--help"):
        print(__doc__)
        return 0

    if tk is None:
        return _report_missing_tkinter()
    if Figure is None:
        return _report_missing_gui_module(
            "matplotlib",
            "The fitted-curve plot needs the Python package 'matplotlib', which is "
            "not available here.\n\n" + str(_MPL_ERR),
        )
    if openpyxl is None:
        return _report_missing_gui_module(
            "openpyxl",
            "Reading the .xlsx workbook needs the Python package 'openpyxl', which is "
            "not available here.",
        )

    root = tk.Tk()
    try:
        root.call("tk", "scaling", 1.25)
    except Exception:
        pass
    app = SpkfpApp(root)
    root.mainloop()
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except SystemExit:
        raise
    except Exception:
        tb = traceback.format_exc()
        print(tb)
        try:
            if tk is not None:
                r = tk.Tk()
                r.withdraw()
                messagebox.showerror("SP-KFP - unexpected error", tb)
                r.destroy()
        except Exception:
            pass
        sys.exit(1)
